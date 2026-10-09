"""Source-file QC and pairing audit only. No model import, fit, prediction or logRR computation."""
from pathlib import Path
import json,hashlib,zipfile,io,math,collections,xml.etree.ElementTree as ET
import pandas as pd
import openpyxl
from docx import Document

root=Path(__file__).resolve().parents[1];ev=root/'evidence';workspace=root.parent
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
frozen={n:sha(workspace/'Schutz_2018_phase0'/n) for n in ['analysis_matrix.tsv','split_manifest.tsv','data_dictionary.tsv','outcome_audit.tsv']}
assert frozen['analysis_matrix.tsv']=='66fd6e7ddd65426225167bfa3016f518cad26ff836e0748decdd3afa39209094'
training=pd.read_csv(workspace/'Schutz_2018_phase0/analysis_matrix.tsv',sep='\t')
studies=pd.read_csv(workspace/'Schutz_2018_phase_minus1/phase_minus1_study_audit.tsv',sep='\t')
assert len(training)==1723 and training.study_id.nunique()==169
predictors=pd.read_csv(workspace/'Schutz_2018_phase0/data_dictionary.tsv',sep='\t').query("role=='main_predictor'").column.tolist()
assert len(predictors)==12
overlap={q:int(studies.astype(str).apply(lambda s:s.str.contains(q,case=False,regex=False)).any(axis=1).sum()) for q in ['2491693','1748102','ME18109','wraf152','Tawaraya','Akyol','Akiyol','Vermeersch','Pannecoucque']}

# Sato: two header-rich supplementary spreadsheets, numbered rows only.
schema={}
for p in (ev/'Sato').glob('*.xlsx'):
 w=openpyxl.load_workbook(p,data_only=True)
 schema[p.name]=[dict(sheet=s.title,rows=s.max_row,columns=s.max_column,header_rows=[list(r) for r in list(s.values)[:6]]) for s in w]
srows=[list(r)[:18] for r in openpyxl.load_workbook(ev/'Sato/TableS2.xlsx',data_only=True).active.values if isinstance(r[0],int)]
keys=[(r[1],r[2],r[3]) for r in srows]
assert len(keys)==len(set(keys))==49
assert all(isinstance(r[j],(int,float)) and math.isfinite(r[j]) and r[j]>0 for r in srows for j in [4,5])
sato=dict(schema=schema,comparisons=49,field_codes=len(set(r[2] for r in srows)),field_years=len(set((r[1],r[2]) for r in srows)),
          counts_by_year=dict(collections.Counter(str(2000+r[1]) for r in srows)),source_reference_counts=dict(collections.Counter(r[17] for r in srows)),
          yield_unit='kg D.W./10 a; multiply both arms by 10 for kg dry weight/ha',unique_pair_keys=49,
          duplicate_full_data_rows=len(srows)-len(set(tuple(r) for r in srows)),n_sd_and_block_ids_public=False,
          unseen_crop='Allium fistulosum; preserve true identity, never alias to garlic',
          strict_predictor_count=2,conditional_predictor_count=1,
          conditional_p_level_mapping={'0':'0 kg P2O5/ha','5':'500 kg P2O5/ha','10':'1000 kg P2O5/ha'},
          p_conversion_factor=0.4364,logRR_computed=False)

# Vermeersch: cultivar-specific absolute grain means from Supplementary Table 2.
d=Document(ev/'Vermeersch_supp_1.docx');assert len(d.tables)==2
arm_rows=[]
for idx,row in enumerate(d.tables[1].rows[2:],3):
 r=[c.text.strip() for c in row.cells]
 for j,cult in enumerate(['RGT Shouna','Lenka','Acardia','Hermes'],4):
  if r[j]=='-':continue
  y=float(r[j]);assert math.isfinite(y) and y>0
  arm_rows.append(dict(source_table=2,source_docx_row=idx,year=int(r[0]),site=r[2],cultivar=cult,strain=r[3],yield_kg_ha_at_15pct_moisture=y))
arm_key=lambda r:(r['year'],r['site'],r['cultivar'],r['strain'])
assert len(arm_rows)==len(set(map(arm_key,arm_rows)))==70
controls={(r['year'],r['site'],r['cultivar']):r for r in arm_rows if r['strain']=='Non-inoculated control'}
pairs=[]
for r in arm_rows:
 if r['strain']=='Non-inoculated control':continue
 k=(r['year'],r['site'],r['cultivar']);c=controls[k]
 pairs.append(dict(experiment_id=f"Vermeersch_{k[0]}_{k[1]}",control_cluster_id=f"Vermeersch_{k[0]}_{k[1]}_{k[2].replace(' ','_')}_uninoculated",
                   year=k[0],site=k[1],cultivar=k[2],strain=r['strain'],treatment_mean_kg_ha_at_15pct_moisture=r['yield_kg_ha_at_15pct_moisture'],
                   control_mean_kg_ha_at_15pct_moisture=c['yield_kg_ha_at_15pct_moisture'],formula='ln(treatment_mean/control_mean); dry conversion 0.85 cancels',
                   treatment_source_docx_row=r['source_docx_row'],control_source_docx_row=c['source_docx_row']))
assert len(pairs)==57 and len(controls)==13

# Köppen grid: original authors' 2006 0.5-degree 1951-2000 climatology, deterministic containing cell.
z=zipfile.ZipFile(ev/'Koeppen-Geiger-ASCII.zip');raw=z.read(next(n for n in z.namelist() if n.lower().endswith('.txt'))).decode('utf8')
grid=[]
for l in raw.splitlines():
 a=l.split()
 try:grid.append((float(a[0]),float(a[1]),a[2]))
 except (IndexError,ValueError):pass
# CODE-ONLY DERIVATION: site observations are input data, not method constants.
# Supply the authentic source metadata; schema and provenance are in docs/INPUT_FORMATS.md.
metadata_path=ev/'external_source_metadata.json'
metadata=json.loads(metadata_path.read_text(encoding='utf8'))
coords=[(r['year'],r['site'],r['latitude'],r['longitude']) for r in metadata['field_sites']]
climates=[]
for year,site,lat,lon in coords:
 matches=[g for g in grid if abs(g[0]-lat)<=0.25 and abs(g[1]-lon)<=0.25]
 assert len(matches)==1,(site,matches)
 glat,glon,code=matches[0]
 climates.append(dict(year=year,site=site,latitude=lat,longitude=lon,grid_latitude=glat,grid_longitude=glon,climate_zone=code))
assert {x['climate_zone'] for x in climates}=={metadata['mapped_predictors']['climate_zone']}

# Retain compatible predictors; leave incompatible/unproven fields missing.
# Preserve the original K2O-to-elemental-K formula; product dose/composition are source inputs.
strict=dict(metadata['mapped_predictors'])
strict['k_applied_kg_ha']=metadata['potassium_product_rate_kg_ha']*metadata['potassium_k2o_mass_fraction']*0.8302
features={k:strict.get(k) for k in predictors}
for p in pairs:p['frozen_predictors']=features.copy()
assert len(strict)==7 and len(set(json.dumps(p['frozen_predictors'],sort_keys=True) for p in pairs))==1
differences=[]
for row in d.tables[0].rows[7:]:
 r=[c.text.strip() for c in row.cells];a=[x['yield_kg_ha_at_15pct_moisture'] for x in arm_rows if x['year']==int(r[0]) and x['site']==r[1] and x['strain']==r[2]]
 if not a:continue
 overall=float(r[7]);avg=sum(a)/len(a)
 differences.append(dict(year=int(r[0]),site=r[1],strain=r[2],table1_overall_mean=overall,table2_unweighted_cultivar_average=avg,difference_kg_ha=overall-avg))
vermeersch=dict(public_docx_tables=2,public_cultivar_specific_arm_means=70,comparisons=57,control_clusters=13,field_years=5,locations=4,
               counts_by_year=dict(collections.Counter(str(r['year']) for r in pairs)),comparisons_by_field_year=dict(collections.Counter(r['experiment_id'] for r in pairs)),
               shared_control_comparison_multiplicity=dict(collections.Counter(collections.Counter(r['control_cluster_id'] for r in pairs).values())),
               strict_predictor_count=7,strict_predictors=strict,missing_predictors=[x for x in predictors if x not in strict],
               conditional_N_P='Table 1 lists only potassium fertilizer; N=P=0 is not assumed in conservative matrix.',
               unique_predictor_profiles=1,climate_mapping=climates,original_design_replicates_per_arm=3,actual_retained_n_sd_public=False,
               source_author_excluded_observations=6,source_author_exclusion_scope_not_recoverable=True,
               table1_table2_summary_disagreement=differences,table1_nonrounding_disagreements=sum(abs(x['difference_kg_ha'])>1 for x in differences),
               outcome_source='Supplementary Table 2 only; Table 1 overall means never pooled with cultivar means.',logRR_computed=False)
out=dict(audit_date='2026-10-06',frozen_hashes=frozen,training_comparisons=len(training),training_studies=training.study_id.nunique(),training_latest_publication_year=int(studies.publication_year.max()),
         frozen_predictors=predictors,training_reference_query_matches=overlap,soybean_training_support=dict(comparisons=int((training.crop=='soybean').sum()),studies=int(training.loc[training.crop=='soybean','study_id'].nunique())),
         soybean_nfix_support=dict(comparisons=len(training.query("crop=='soybean' and inoculant_function==0")),studies=int(training.query("crop=='soybean' and inoculant_function==0").study_id.nunique())),
         sato=sato,vermeersch=vermeersch,model_fits=0,model_predictions=0,model_performance_files_read=0)
(root/'structural_audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
(root/'vermeersch_pairing_audit.json').write_text(json.dumps(pairs,ensure_ascii=False,indent=2),encoding='utf8')
(root/'source_hashes.json').write_text(json.dumps({str(p.relative_to(root)):sha(p) for p in ev.rglob('*') if p.is_file()},ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(dict(sato_pairs=49,sato_field_years=sato['field_years'],vermeersch_pairs=57,vermeersch_field_years=5,vermeersch_predictors=len(strict),climates=climates,frozen_hashes=frozen,table_disagreements=vermeersch['table1_nonrounding_disagreements']),ensure_ascii=True))
